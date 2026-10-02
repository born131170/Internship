from __future__ import annotations
import json, shutil, subprocess, threading, time, traceback, uuid
from pathlib import Path
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse
import pipeline, snapshots, llm
import scoring

DATA=Path("data"); SNAP=DATA/"snapshots"; STATIC=Path(__file__).parent/"static"
DATA.mkdir(exist_ok=True); SNAP.mkdir(parents=True, exist_ok=True)
JOBS: dict = {}
LLM_JOBS: dict = {}
app=FastAPI(title="PersonaScope")

@app.middleware("http")
async def _nocache(request, call_next):
    resp=await call_next(request)
    if request.url.path=="/" or request.url.path.startswith("/api/"):
        resp.headers["Cache-Control"]="no-store, must-revalidate"
    return resp

def _vdir(vid: str) -> Path:
    d=DATA/vid
    if not d.is_dir(): raise HTTPException(404,"видео не найдено")
    return d

def _metrics(d: Path):
    p=d/"metrics.jsonl"
    if not p.exists(): raise HTTPException(409,"анализ ещё не выполнен")
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]

def _clean(o):
    if isinstance(o,float): return None if o!=o else o
    if isinstance(o,dict): return {k:_clean(v) for k,v in o.items()}
    if isinstance(o,list): return [_clean(v) for v in o]
    return o

def _ffmpeg_exe():
    try:
        import imageio_ffmpeg; return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        import shutil as _sh; return _sh.which("ffmpeg")

def _transcode(d: Path, which="dashboard"):
    """mp4v/HEVC не читаются браузерами -> H.264 *_web.mp4."""
    if which=="dashboard": src,dst,extra=d/"dashboard.mp4",d/"dashboard_web.mp4",()
    else: src,dst,extra=d/"source.mp4",d/"source_web.mp4",("-vf","scale='min(1280,iw)':-2")
    if not src.exists() or dst.exists(): return dst.exists()
    exe=_ffmpeg_exe()
    if not exe:
        print("[WARN] ffmpeg не найден: pip install imageio-ffmpeg")
        return False
    try:
        subprocess.run([exe,"-y","-i",str(src),*extra,"-c:v","libx264","-preset","veryfast","-crf","23",
                        "-pix_fmt","yuv420p","-movflags","+faststart","-an",str(dst)],
                       check=True,capture_output=True,timeout=3600)
        return dst.exists()
    except Exception:
        traceback.print_exc(); return False

def _register(vid: str):
    reg=DATA/"index.json"; items=[]
    if reg.exists():
        try: items=json.loads(reg.read_text(encoding="utf-8"))
        except Exception: items=[]
    items=[x for x in items if x.get("id")!=vid]
    items.append({"id":vid,"finished":time.time()})
    reg.write_text(json.dumps(items,ensure_ascii=False),encoding="utf-8")

@app.get("/")
def index(): return FileResponse(STATIC/"index.html")

@app.post("/api/videos")
async def upload(file: UploadFile=File(...)):
    vid=uuid.uuid4().hex[:10]; d=DATA/vid; d.mkdir(parents=True)
    with (d/"source.mp4").open("wb") as f: shutil.copyfileobj(file.file,f)
    return {"id":vid,"name":file.filename}

@app.post("/api/videos/{vid}/analyze")
def analyze(vid: str, stride: int=2, width: int=640, clips: bool=True):
    d=_vdir(vid); JOBS[vid]={"state":"running","progress":0.0,"msg":"старт"}
    def run():
        try:
            pipeline.analyze_video(d/"source.mp4", d, stride=stride, width=width, make_clips=clips,
                progress_cb=lambda f,m: JOBS[vid].update(progress=f,msg=m))
            _transcode(d)
            _register(vid)
            JOBS[vid].update(state="done",progress=1.0,msg="готово")
        except Exception as e:
            traceback.print_exc()
            msg = str(e) if isinstance(e, RuntimeError) else repr(e)
            JOBS[vid].update(state="error", msg=msg)
    threading.Thread(target=run,daemon=True).start()
    return {"job":vid}

@app.get("/api/jobs/{vid}")
def job(vid: str): return JOBS.get(vid,{"state":"idle","progress":0.0,"msg":""})

@app.get("/api/videos/latest")
def latest():
    reg=DATA/"index.json"
    if reg.exists():
        try:
            items=json.loads(reg.read_text(encoding="utf-8"))
            if items: return {"id":items[-1]["id"],"finished":items[-1].get("finished")}
        except Exception: pass
    cands=[((p/"summary.json").stat().st_mtime,p) for p in DATA.glob("*") if (p/"summary.json").exists()]
    if not cands: return {"id":None}
    cands.sort(); return {"id":cands[-1][1].name}

@app.get("/api/videos/{vid}/summary")
def summary(vid: str):
    p=_vdir(vid)/"summary.json"
    if not p.exists(): raise HTTPException(409,"анализ ещё не выполнен")
    return _clean(json.loads(p.read_text(encoding="utf-8")))

@app.get("/api/videos/{vid}/metrics")
def metrics(vid: str): return _clean(_metrics(_vdir(vid)))

@app.get("/api/videos/{vid}/truth")
def truth(vid: str):
    s=json.loads((_vdir(vid)/"summary.json").read_text(encoding="utf-8"))
    return {"cues":_clean(s["truth_cues"]),"heuristic":s["truth_heuristic"]}

@app.get("/api/videos/{vid}/file")
def vfile(vid: str, path: str):
    base=_vdir(vid).resolve(); fp=(base/path).resolve()
    if not str(fp).startswith(str(base)): raise HTTPException(403)
    if not fp.exists():
        if fp.name in ("dashboard_web.mp4","source_web.mp4"):
            w="dashboard" if fp.name.startswith("dashboard") else "source"
            if not _transcode(base,w): raise HTTPException(500,"H.264-версия недоступна: pip install imageio-ffmpeg и перезапустите сервер")
        else: raise HTTPException(404)
    return FileResponse(fp, media_type="video/mp4" if fp.suffix==".mp4" else None, headers={"Cache-Control":"no-store"})

@app.post("/api/videos/{vid}/snapshots")
def create_snap(vid: str, payload: dict):
    try:
        d=_vdir(vid); recs=_metrics(d)
        if not recs: raise ValueError("нет кадровых метрик видео")
        t0=float(payload.get("t0",0)); t1=float(payload.get("t1",0))
        if t1<=t0: raise ValueError("конец окна должен быть позже начала")
        t0=max(t0,recs[0]["t"]); t1=min(t1,recs[-1]["t"])
        if t1-t0<0.3: raise ValueError("окно короче 0.3 c — расширьте выделение")
        snap=snapshots.create_snapshot(recs,t0,t1,
            payload.get("channels") or snapshots.DEFAULT_CHANNELS, payload.get("name") or "слепок")
        snap["source_video"]=vid
        (SNAP/f"{snap['id']}.json").write_text(json.dumps(_clean(snap),ensure_ascii=False),encoding="utf-8")
        return _clean(snap)
    except ValueError as e:
        return JSONResponse(status_code=422,content={"error":f"Некорректное окно слепка: {e}"})
    except Exception as e:
        traceback.print_exc()
        return JSONResponse(status_code=502,content={"error":f"Ошибка слепка: {llm.describe_error(e)}"})

@app.get("/api/snapshots")
def list_snaps():
    out=[]
    for p in sorted(SNAP.glob("*.json")):
        s=json.loads(p.read_text(encoding="utf-8")); s.pop("series",None); out.append(s)
    return out

@app.get("/api/snapshots/{sid}")
def get_snap(sid: str):
    p=SNAP/f"{sid}.json"
    if not p.exists(): raise HTTPException(404)
    return json.loads(p.read_text(encoding="utf-8"))

@app.delete("/api/snapshots/{sid}")
def del_snap(sid: str):
    p=SNAP/f"{sid}.json"
    if not p.exists(): raise HTTPException(404,"слепок не найден")
    p.unlink(); return {"ok":True}

@app.post("/api/snapshots/upload")
async def upload_snap(file: UploadFile=File(...)):
    data=json.loads(await file.read()); data.setdefault("id",uuid.uuid4().hex[:10])
    (SNAP/f"{data['id']}.json").write_text(json.dumps(data,ensure_ascii=False),encoding="utf-8")
    return {"id":data["id"]}

@app.post("/api/videos/{vid}/search")
def search(vid: str, payload: dict):
    try:
        recs=_metrics(_vdir(vid)); sid=payload.get("snapshot_id")
        if sid: snap=json.loads((SNAP/f"{sid}.json").read_text(encoding="utf-8"))
        elif payload.get("snapshot"): snap=payload["snapshot"]
        else: return JSONResponse(status_code=422,content={"error":"нужен snapshot_id или snapshot"})
        mode=payload.get("mode") or ("episodes" if snap.get("pattern") else "free")
        res,meta=snapshots.search(recs,snap,hop=float(payload.get("hop",0.25)),
                                  top_k=int(payload.get("top_k",10)),mode=mode)
        return {"snapshot":snap["id"],"results":res,"total":len(res),"meta":meta}
    except Exception as e:
        traceback.print_exc()
        return JSONResponse(status_code=502,content={"error":f"Ошибка поиска: {llm.describe_error(e)}"})

@app.get("/api/llm/config")
def get_cfg(): return llm.load_config(DATA)

@app.put("/api/llm/config")
def put_cfg(payload: dict): llm.save_config(DATA,payload); return payload

@app.get("/api/llm/status")
def llm_status():
    cfg=llm.load_config(DATA)
    return {"configured":bool(cfg.get("base_url") and cfg.get("api_key") and cfg.get("model")),
            "base_url":cfg.get("base_url"),"model":cfg.get("model"),"key_set":bool(cfg.get("api_key"))}

@app.post("/api/llm/test")
def test_cfg(payload: dict = None):
    cfg=payload if isinstance(payload,dict) and payload.get("base_url") else llm.load_config(DATA)
    if not cfg.get("api_key"):
        return JSONResponse(status_code=409,content={"ok":False,"error":"API key не задан"})
    try:
        return {"ok":True,"reply":llm.ping(cfg)}
    except Exception as e:
        return JSONResponse(status_code=502,content={"ok":False,"error":llm.describe_error(e)})

@app.post("/api/videos/{vid}/llm/analyze")
def llm_analyze(vid: str, payload: dict = None):
    """Фоновый LLM-анализ: мгновенный старт + опрос статуса (не зависит от таймаута прокси)."""
    payload=payload or {}
    try:
        d=_vdir(vid)
        s=json.loads((d/"summary.json").read_text(encoding="utf-8"))
    except Exception as e:
        traceback.print_exc()
        return JSONResponse(status_code=404,content={"error":f"Нет данных анализа для видео {vid}: {e}"})
    cfg=llm.load_config(DATA)
    if not (cfg.get("base_url") and cfg.get("api_key") and cfg.get("model")):
        return JSONResponse(status_code=409,content={"error":"LLM не настроен: укажите Base URL, API key и модель в «Настройки LLM» и нажмите «Сохранить»"})
    st=LLM_JOBS.get(vid)
    if st and st.get("state")=="running":
        return {"state":"running","msg":st.get("msg","")}
    def run():
        st.update(state="running",msg="запрос к LLM…")
        try:
            scores=scoring.compute_scores(s)
            warnings=[]; evidence=llm.build_evidence(s); raw=None; finish=None
            evidence["scores_block"]=scores
            for cap in (4,2,0):
                try:
                    raw,finish=llm.run_analysis(cfg,payload.get("prompt") or llm.DEFAULT_USER_PROMPT,evidence,cite_cap=cap)
                    break
                except Exception as e:
                    if cap==0 or not llm.is_retriable(e):
                        traceback.print_exc()
                        st.update(state="error",error=llm.describe_error(e),msg="ошибка"); return
                    warnings.append(f"API: {llm.describe_error(e)} — повторный запрос с ограничением доказательств (≤{2 if cap==4 else 0} на систему)")
                    st.update(msg=f"повтор запроса (лимит доказательств ≤{2 if cap==4 else 0})…")
            if finish=="length":
                warnings.append("Ответ модели оборван на лимите токенов прокси (finish_reason=length)")
            parsed=llm.parse_json(raw)
            if parsed is None:
                fixed=llm._repair_json((raw or "").strip())
                if fixed is not None:
                    parsed=fixed; warnings.append("Локальная починка: оборванный JSON восстановлен закрытием скобок")
            if parsed is None:
                st.update(msg="авторемонт JSON…")
                try:
                    raw2=llm.repair_json_response(cfg,raw or "")
                    parsed2=llm.parse_json(raw2)
                    if parsed2 is not None:
                        raw=raw2; parsed=parsed2; warnings.append("Авторемонт: повторный запрос вернул валидный JSON")
                except Exception as e:
                    warnings.append(f"Авторемонт не удался: {llm.describe_error(e)}")
            if parsed is None:
                warnings.append("Ответ модели не распарсен как JSON — результаты НЕ применены, смотрите RAW")
            else:
                try:
                    parsed,w=llm.validate_result(parsed,s); warnings+=w
                except Exception as e:
                    warnings.append(f"Валидация ответа: {e}")
                try:
                    parsed=scoring.enrich_parsed(parsed,s)
                except Exception as e:
                    warnings.append(f"Автодоказательства: {e}")
            result={"raw":raw,"parsed":parsed,"warnings":warnings,"scores":scores}
            (d/"llm_result.json").write_text(json.dumps(result,ensure_ascii=False),encoding="utf-8")
            st.update(state="done",result=result,msg="готово")
        except Exception as e:
            traceback.print_exc()
            st.update(state="error",error=f"Внутренняя ошибка сервера: {llm.describe_error(e)}",msg="ошибка")
    st={"state":"queued","msg":"подготовка запроса…","result":None,"error":None,"thread":None}
    LLM_JOBS[vid]=st
    t=threading.Thread(target=run,daemon=True); st["thread"]=t; t.start()
    return {"state":"running","started":True}

@app.get("/api/videos/{vid}/llm/status")
def llm_analyze_status(vid: str):
    st=LLM_JOBS.get(vid)
    if not st:
        d=_vdir(vid)
        if (d/"llm_result.json").exists():
            return {"state":"done","result":json.loads((d/"llm_result.json").read_text(encoding="utf-8")),"msg":"готово (с диска)"}
        return {"state":"idle","msg":""}
    out={"state":st["state"],"msg":st.get("msg","")}
    if st["state"]=="done": out["result"]=st["result"]
    if st["state"]=="error": out["error"]=st.get("error")
    return out

SYSREP=[("big_five","Big Five",lambda p:[("Openness",p.get("openness")),("Conscientiousness",p.get("conscientiousness")),("Extraversion",p.get("extraversion")),("Agreeableness",p.get("agreeableness")),("Neuroticism",p.get("neuroticism"))]),
 ("mbti","MBTI",lambda p:[("E/I",(p.get("axes") or {}).get("E_I")),("S/N",(p.get("axes") or {}).get("S_N")),("T/F",(p.get("axes") or {}).get("T_F")),("J/P",(p.get("axes") or {}).get("J_P")),("type",p.get("type"))]),
 ("enneagram","Эннеаграмма",lambda p:[("type",p.get("type")),("wing",p.get("wing"))]),
 ("temperament","Temperament",lambda p:[("Сангвиник",p.get("sanguine")),("Холерик",p.get("choleric")),("Меланхолик",p.get("melancholic")),("Флегматик",p.get("phlegmatic"))]),
 ("hexaco","HEXACO",lambda p:[("H",p.get("H")),("E",p.get("E")),("X",p.get("X")),("A",p.get("A")),("C",p.get("C")),("O",p.get("O"))]),
 ("pid5","PID-5 (DSM-5)",lambda p:[("NegAffect",p.get("negative_affect")),("Detachment",p.get("detachment")),("Antagonism",p.get("antagonism")),("Disinhibition",p.get("disinhibition")),("Psychoticism",p.get("psychoticism"))])]

@app.get("/api/videos/{vid}/report/pptx")
def report_pptx(vid: str):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception:
        return JSONResponse(status_code=502,content={"error":"не установлен python-pptx: pip install python-pptx"})
    d=_vdir(vid)
    s=json.loads((d/"summary.json").read_text(encoding="utf-8"))
    res=json.loads((d/"llm_result.json").read_text(encoding="utf-8")) if (d/"llm_result.json").exists() else None
    prs=Presentation(); blank=prs.slide_layouts[6]
    def slide(t):
        sl=prs.slides.add_slide(blank)
        tb=sl.shapes.add_textbox(Inches(0.5),Inches(0.3),Inches(9),Inches(0.9)); tf=tb.text_frame
        tf.text=t; tf.paragraphs[0].font.size=Pt(26); tf.paragraphs[0].font.bold=True
        return sl
    def body(sl,text,size=14):
        tb=sl.shapes.add_textbox(Inches(0.5),Inches(1.2),Inches(9),Inches(5.6)); tf=tb.text_frame; tf.word_wrap=True
        for i,ln in enumerate(text.split("\n")):
            pp=tf.paragraphs[0] if i==0 else tf.add_paragraph(); pp.text=ln; pp.font.size=Pt(size)
    sl=slide("PersonaScope — отчёт по анализу личности")
    body(sl,f"Видео: {s.get('video')}\nДлительность: {s.get('duration_sec')} с · кадров: {s.get('n_frames')}\nСформирован: {time.strftime('%Y-%m-%d %H:%M')}")
    sl=slide("Индикаторы MediaPipe (кадров / событий)")
    body(sl,"\n".join(f"{k} {s.get('pattern_names',{}).get(k,k)}: {v} кадров, {s.get('events',{}).get(k,0)} событий" for k,v in s.get("counters",{}).items()))
    p=(res or {}).get("parsed")
    if not p:
        sl=slide("LLM-анализ не выполнялся"); body(sl,"Разделы оценок пусты: запустите LLM-анализ в интерфейсе.")
    else:
        sl=slide("Заключение"); body(sl,(p.get("summary") or "")+f"\nconfidence: {p.get('confidence')}")
        eps={e["id"]:e for e in s.get("episodes",[])}
        for key,title,rows in SYSREP:
            dblk=p.get(key) or {}
            sl=slide(title)
            lines=[f"{n}: {('—' if v is None else round(v,1) if isinstance(v,(int,float)) else v)}" for n,v in rows(dblk)]
            if dblk.get("notes"): lines.append("Примечание: "+str(dblk.get("notes")))
            ev=(p.get("evidence") or {}).get(key) or []
            lines.append("Доказательства: "+str(len(ev)))
            body(sl,"\n".join(lines))
            for i,e in enumerate(ev[:6]):
                ep=eps.get(e.get("episode_id"))
                if ep and (d/ep["thumb"]).exists():
                    sl.shapes.add_picture(str(d/ep["thumb"]),Inches(0.5+(i%3)*3.1),Inches(3.0+(i//3)*2.0),width=Inches(2.9))
        t=p.get("truthfulness") or {}
        sl=slide("Правдивость"); body(sl,f"Score: {t.get('score')}\nВердикт: {t.get('verdict')}\n"+
            "\n".join(f"- {c.get('cue')} ({c.get('direction')})" for c in (t.get("cues") or [])))
    out=d/"report.pptx"; prs.save(str(out))
    return FileResponse(out,filename=f"personascope_{vid}.pptx",
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")

@app.get("/api/videos/{vid}/scores")
def scores_ep(vid: str):
    s=json.loads((_vdir(vid)/"summary.json").read_text(encoding="utf-8"))
    return scoring.compute_scores(s)

@app.get("/api/videos/{vid}/llm/result")
def llm_result(vid: str):
    p=_vdir(vid)/"llm_result.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None
